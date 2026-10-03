"""Credential-free tests for protocol registration and adapter guardrails."""

from __future__ import annotations

import asyncio
import contextlib
import io
import json
import logging
import subprocess
import unittest
import zipfile

import anyio
import httpx
from mcp import ClientSession
from mcp.shared.message import SessionMessage
from synology_file_station_mcp.server import (
    AdapterError,
    FileStationClient,
    Limits,
    _check_zip_limits,
    _convert_document,
    _decode_path,
    _error_result,
    _ffprobe,
    _image_content,
    _pdf_page,
    build_server,
)


class FakeAdapter:
    """Provide deterministic tool responses for MCP registration tests."""

    def list_entries(self, path, page, page_size):
        return {"path": path, "page": page, "page_size": page_size}

    def search(self, path, pattern, page_size, page=1):
        return {"path": path, "pattern": pattern, "page_size": page_size, "page": page}

    def stat(self, path):
        return {"path": path}

    def read_text(self, path):
        return {"path": path, "text": "fixture"}

    def read_document(self, path):
        return {"path": path, "markdown": "fixture"}

    def preview_image(self, path):
        return {"path": path}

    def preview_pdf_page(self, path, page):
        return {"path": path, "page": page}

    def media_info(self, path):
        return {"path": path, "inspection_limited": True}


class ServerProtocolTests(unittest.IsolatedAsyncioTestCase):
    """Exercise a real MCP SDK v2 initialize/list/call session."""

    async def test_initialize_list_tools_and_call(self):
        server = build_server(FakeAdapter())
        client_to_server_send, client_to_server_receive = (
            anyio.create_memory_object_stream[SessionMessage | Exception](100)
        )
        server_to_client_send, server_to_client_receive = (
            anyio.create_memory_object_stream[SessionMessage](100)
        )
        server_task = asyncio.create_task(
            server._lowlevel_server.run(
                client_to_server_receive,
                server_to_client_send,
                server._lowlevel_server.create_initialization_options(),
            )
        )
        try:
            async with ClientSession(
                server_to_client_receive, client_to_server_send
            ) as session:
                initialized = await session.initialize()
                self.assertEqual(
                    initialized.server_info.name, "synology-file-station-mcp"
                )
                listed = await session.list_tools()
                names = {tool.name for tool in listed.tools}
                self.assertEqual(
                    names,
                    {
                        "list_entries",
                        "search_by_name",
                        "stat",
                        "read_text",
                        "read_document",
                        "preview_image",
                        "preview_pdf_page",
                        "media_info",
                    },
                )
                called = await session.call_tool(
                    "list_entries", {"path": ".", "page": 1, "page_size": 100}
                )
                self.assertFalse(called.is_error)
                self.assertIn('"path": "."', called.content[0].text)
        finally:
            await client_to_server_send.aclose()
            await server_to_client_receive.aclose()
            server_task.cancel()
            with contextlib.suppress(asyncio.CancelledError):
                await server_task


class PathAndErrorTests(unittest.TestCase):
    """Cover traversal, archive limits and secret-safe error rendering."""

    def test_path_rejects_traversal_and_encoded_separators(self):
        for path in (
            "../secret",
            "a/../secret",
            "/absolute",
            "a%2f..%2fsecret",
            "a\\b",
            "a%00b",
        ):
            with self.subTest(path=path), self.assertRaises(AdapterError):
                _decode_path(path)

    def test_archive_expansion_limit(self):
        payload = io.BytesIO()
        import zipfile

        with zipfile.ZipFile(payload, "w") as archive:
            archive.writestr("safe.txt", "x" * 32)
        _check_zip_limits(payload.getvalue(), Limits(archive_expanded_bytes=64))
        with zipfile.ZipFile(payload, "w") as archive:
            archive.writestr("safe.txt", "x" * 65)
        with self.assertRaises(AdapterError):
            _check_zip_limits(payload.getvalue(), Limits(archive_expanded_bytes=64))

    def test_unexpected_error_does_not_expose_secret(self):
        with self.assertLogs(
            "synology_file_station_mcp", level=logging.ERROR
        ) as captured:
            result = _error_result(RuntimeError("password=super-secret"))
        self.assertTrue(result.is_error)
        self.assertNotIn("super-secret", "".join(captured.output))
        self.assertNotIn("super-secret", result.content[0].text)

    def test_native_image_content(self):
        from PIL import Image

        image = Image.new("RGB", (3, 2), (12, 34, 56))
        payload = io.BytesIO()
        image.save(payload, format="PNG")
        content = _image_content(payload.getvalue(), Limits())
        self.assertEqual(content.type, "image")
        self.assertEqual(content.mime_type, "image/png")
        self.assertGreater(len(content.data), 20)

    def test_local_document_converters(self):
        limits = Limits(output_bytes=2 * 1024 * 1024)
        fixtures = {
            ".docx": _zip_fixture(
                {
                    "[Content_Types].xml": "<?xml version='1.0'?><Types xmlns='http://schemas.openxmlformats.org/package/2006/content-types'><Default Extension='rels' ContentType='application/vnd.openxmlformats-package.relationships+xml'/><Default Extension='xml' ContentType='application/xml'/><Override PartName='/word/document.xml' ContentType='application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml'/></Types>",
                    "word/document.xml": "<?xml version='1.0'?><w:document xmlns:w='http://schemas.openxmlformats.org/wordprocessingml/2006/main'><w:body><w:p><w:r><w:t>fixture docx</w:t></w:r></w:p></w:body></w:document>",
                }
            ),
            ".xlsx": _xlsx_fixture(),
            ".pptx": _pptx_fixture(),
        }
        for extension, fixture in fixtures.items():
            with self.subTest(extension=extension):
                result = _convert_document(fixture, f"fixture{extension}", limits)
                self.assertIsInstance(result, str)

    def test_local_pdf_page_preview(self):
        page = _pdf_page(_minimal_pdf(), 1, Limits())
        self.assertEqual(page.type, "image")
        self.assertEqual(page.mime_type, "image/png")

    def test_synthetic_audio_and_video_probe_is_allowlisted(self):
        import tempfile

        for kind, suffix, expected_type in (
            ("audio", ".aac", "audio"),
            ("video", ".mp4", "video"),
        ):
            with (
                self.subTest(kind=kind),
                tempfile.NamedTemporaryFile(suffix=suffix) as temporary,
            ):
                temporary.write(_synthetic_media(kind))
                temporary.flush()
                probe = _ffprobe(temporary.name)
            self.assertIsNotNone(probe)
            self.assertTrue(
                any(
                    stream.get("codec_type") == expected_type
                    for stream in probe.get("streams", [])
                )
            )
            self.assertNotIn("tags", json.dumps(probe))


class HttpGuardrailTests(unittest.TestCase):
    """Exercise redirect, authentication and range validation without a NAS."""

    @staticmethod
    def _method(request: httpx.Request) -> str | None:
        """Read File Station's documented POST form or GET query method."""

        if request.method == "POST":
            return httpx.QueryParams(request.content.decode("ascii")).get("method")
        return request.url.params.get("method")

    def _client(self, handler):
        transport = httpx.MockTransport(handler)
        return FileStationClient(
            username="fixture-user",
            password="fixture-password",
            http_client=httpx.Client(
                base_url="https://nas.bohdal.name:5001",
                transport=transport,
                follow_redirects=False,
                trust_env=False,
            ),
            limits=Limits(page_size=10),
        )

    def test_redirect_is_rejected(self):
        def handler(request):
            return httpx.Response(
                302, headers={"location": "https://evil.example/"}, request=request
            )

        client = self._client(handler)
        with self.assertRaisesRegex(AdapterError, "Redirects"):
            client._login()
        client.close()

    def test_auth_failure_has_no_password_in_message(self):
        def handler(request):
            return httpx.Response(
                200, json={"success": False, "error": {"code": 400}}, request=request
            )

        client = self._client(handler)
        with self.assertRaises(AdapterError) as raised:
            client._login()
        self.assertNotIn("fixture-password", str(raised.exception))
        client.close()

    def test_media_root_requires_real_path_match(self):
        def handler(request):
            if request.url.path.endswith("auth.cgi"):
                return httpx.Response(
                    200,
                    json={"success": True, "data": {"sid": "memory-sid"}},
                    request=request,
                )
            return httpx.Response(
                200,
                json={
                    "success": True,
                    "data": {
                        "shares": [
                            {
                                "name": "Media",
                                "path": "/Media",
                                "additional": {
                                    "real_path": "/volume1/Other",
                                    "mount_point_type": "local",
                                },
                            }
                        ]
                    },
                },
                request=request,
            )

        client = self._client(handler)
        with self.assertRaisesRegex(AdapterError, "not verified"):
            client.verify_root()
        client.close()

    def test_documented_share_and_listing_shapes_are_verified(self):
        def handler(request):
            if request.url.path.endswith("auth.cgi"):
                return httpx.Response(
                    200,
                    json={"success": True, "data": {"sid": "memory-sid"}},
                    request=request,
                )
            method = self._method(request)
            if method == "list_share":
                return httpx.Response(
                    200,
                    json={
                        "success": True,
                        "data": {
                            "shares": [
                                {
                                    "name": "Media",
                                    "path": "/Media",
                                    "additional": {
                                        "real_path": "/volume1/Media",
                                        "mount_point_type": "local",
                                    },
                                }
                            ]
                        },
                    },
                    request=request,
                )
            self.assertEqual(method, "list")
            return httpx.Response(
                200,
                json={
                    "success": True,
                    "data": {
                        "files": [
                            {
                                "name": "guide.txt",
                                "path": "/Media/guide.txt",
                                "isdir": False,
                                "additional": {
                                    "real_path": "/volume1/Media/guide.txt",
                                    "size": 7,
                                    "time": 1,
                                    "mount_point_type": "local",
                                },
                            }
                        ]
                    },
                },
                request=request,
            )

        client = self._client(handler)
        result = client.list_entries(".", 1, 10)
        self.assertEqual(result["items"][0]["path"], "guide.txt")
        client.close()

    def test_large_media_requires_verified_seekable_range(self):
        methods = []
        media = _synthetic_media("video")
        self.assertGreater(len(media), 10)

        def handler(request):
            if request.url.path.endswith("auth.cgi"):
                return httpx.Response(
                    200,
                    json={"success": True, "data": {"sid": "memory-sid"}},
                    request=request,
                )
            method = self._method(request)
            methods.append((method, request.headers.get("range")))
            if method == "list_share":
                return httpx.Response(
                    200,
                    json={
                        "success": True,
                        "data": {
                            "shares": [
                                {
                                    "name": "Media",
                                    "path": "/Media",
                                    "additional": {
                                        "real_path": "/volume1/Media",
                                        "mount_point_type": "local",
                                    },
                                }
                            ]
                        },
                    },
                    request=request,
                )
            if method == "getinfo":
                return httpx.Response(
                    200,
                    json={
                        "success": True,
                        "data": {
                            "files": [
                                {
                                    "name": "clip.mp4",
                                    "path": "/Media/clip.mp4",
                                    "isdir": False,
                                    "additional": {
                                        "real_path": "/volume1/Media/clip.mp4",
                                        "size": len(media),
                                        "type": "mp4",
                                        "mount_point_type": "local",
                                    },
                                }
                            ]
                        },
                    },
                    request=request,
                )
            self.assertEqual(method, "download")
            range_header = request.headers.get("range")
            self.assertIsNotNone(range_header)
            end = len(media) - 1
            self.assertEqual(range_header, f"bytes=0-{end}")
            return httpx.Response(
                206,
                headers={
                    "content-range": f"bytes 0-{end}/{len(media)}",
                    "content-length": str(len(media)),
                },
                content=media,
                request=request,
            )

        client = self._client(handler)
        client.limits = Limits(
            page_size=10, media_small_bytes=10, media_probe_bytes=64 * 1024
        )
        result = client.media_info("clip.mp4")
        self.assertFalse(result["inspection_limited"])
        self.assertIn("ffprobe", result)
        self.assertTrue(
            any(method == "download" and range_value for method, range_value in methods)
        )
        client.close()

    def test_missing_size_returns_basic_metadata_without_download(self):
        def handler(request):
            if request.url.path.endswith("auth.cgi"):
                return httpx.Response(
                    200,
                    json={"success": True, "data": {"sid": "memory-sid"}},
                    request=request,
                )
            method = self._method(request)
            if method == "list_share":
                return httpx.Response(
                    200,
                    json={
                        "success": True,
                        "data": {
                            "shares": [
                                {
                                    "name": "Media",
                                    "path": "/Media",
                                    "additional": {
                                        "real_path": "/volume1/Media",
                                        "mount_point_type": "local",
                                    },
                                }
                            ]
                        },
                    },
                    request=request,
                )
            return httpx.Response(
                200,
                json={
                    "success": True,
                    "data": {
                        "files": [
                            {
                                "name": "clip.mp4",
                                "path": "/Media/clip.mp4",
                                "isdir": False,
                                "additional": {
                                    "real_path": "/volume1/Media/clip.mp4",
                                    "mount_point_type": "local",
                                },
                            }
                        ]
                    },
                },
                request=request,
            )

        client = self._client(handler)
        result = client.media_info("clip.mp4")
        self.assertTrue(result["inspection_limited"])
        self.assertIn("file size", result["limitation"])
        client.close()

    def test_search_cleanup_runs_after_completion(self):
        methods = []

        def handler(request):
            if request.url.path.endswith("auth.cgi"):
                return httpx.Response(
                    200,
                    json={"success": True, "data": {"sid": "memory-sid"}},
                    request=request,
                )
            method = self._method(request)
            methods.append(method)
            if method == "list_share":
                return httpx.Response(
                    200,
                    json={
                        "success": True,
                        "data": {
                            "shares": [
                                {
                                    "name": "Media",
                                    "path": "/Media",
                                    "additional": {
                                        "real_path": "/volume1/Media",
                                        "mount_point_type": "local",
                                    },
                                }
                            ]
                        },
                    },
                    request=request,
                )
            if method == "start":
                return httpx.Response(
                    200,
                    json={"success": True, "data": {"taskid": "fixture-task"}},
                    request=request,
                )
            if method == "list":
                return httpx.Response(
                    200,
                    json={"success": True, "data": {"finished": True, "files": []}},
                    request=request,
                )
            return httpx.Response(
                200, json={"success": True, "data": {}}, request=request
            )

        client = self._client(handler)
        result = client.search(".", "clip", 10)
        self.assertEqual(result["items"], [])
        self.assertEqual(methods[-2:], ["stop", "clean"])
        client.close()


def _zip_fixture(files):
    """Build a tiny office fixture without credentials or external files."""

    payload = io.BytesIO()
    with zipfile.ZipFile(payload, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for name, content in files.items():
            archive.writestr(name, content)
    return payload.getvalue()


def _xlsx_fixture():
    """Create a valid XLSX fixture through the pinned parser dependency."""

    from openpyxl import Workbook

    workbook = Workbook()
    workbook.active["A1"] = "fixture xlsx"
    payload = io.BytesIO()
    workbook.save(payload)
    return payload.getvalue()


def _pptx_fixture():
    """Create a valid PPTX fixture through the pinned parser dependency."""

    from pptx import Presentation

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "fixture pptx"
    slide.placeholders[1].text = "MCP test"
    payload = io.BytesIO()
    presentation.save(payload)
    return payload.getvalue()


def _minimal_pdf():
    """Return a small valid one-page PDF for the native preview test."""

    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 100 100] /Contents 4 0 R /Resources << >> >>",
        b"<< /Length 0 >>\nstream\n\nendstream",
    ]
    body = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for index, value in enumerate(objects, start=1):
        offsets.append(len(body))
        body.extend(f"{index} 0 obj\n".encode())
        body.extend(value)
        body.extend(b"\nendobj\n")
    xref = len(body)
    body.extend(f"xref\n0 {len(objects) + 1}\n".encode())
    body.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        body.extend(f"{offset:010d} 00000 n \n".encode())
    body.extend(
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    )
    return bytes(body)


def _synthetic_media(kind):
    """Create a tiny credential-free media fixture through local ffmpeg."""

    source = (
        "color=c=black:s=32x32:d=1" if kind == "video" else "anullsrc=r=8000:cl=mono"
    )
    input_args = ["-f", "lavfi", "-i", source, "-t", "1"]
    output_args = (
        ["-f", "mp4", "-movflags", "frag_keyframe+empty_moov", "pipe:1"]
        if kind == "video"
        else ["-f", "adts", "-c:a", "aac", "pipe:1"]
    )
    result = subprocess.run(
        ["ffmpeg", "-loglevel", "error", *input_args, *output_args],
        capture_output=True,
        check=True,
        timeout=10,
    )
    return result.stdout


if __name__ == "__main__":
    unittest.main()
